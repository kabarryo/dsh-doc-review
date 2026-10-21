#!/usr/bin/env python3
"""Tests for scripts/xlsx_edit.py.

Run from the repository root:
  python3 -m unittest discover -s tests -v
"""
from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from make_fixtures import build_rich, build_simple  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
HELPER = ROOT / "scripts" / "xlsx_edit.py"


class HelperCase(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="xlsx-edit-"))
        self.home = self.tmp / "home"
        self.home.mkdir()
        self.rich = build_rich(self.tmp / "rich.xlsx")
        self.simple = build_simple(self.tmp / "simple.xlsx")

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    # -- helpers ---------------------------------------------------------
    def run_helper(self, file: Path, ops, *args: str):
        env = dict(os.environ, DOC_REVIEW_HOME=str(self.home))
        command = [
            sys.executable,
            str(HELPER),
            "--file",
            str(file),
            "--ops",
            json.dumps(ops, ensure_ascii=False),
            *args,
        ]
        return subprocess.run(command, capture_output=True, text=True, env=env)

    def digest(self, path: Path) -> str:
        return hashlib.sha256(Path(path).read_bytes()).hexdigest()

    def parts(self, path: Path) -> set:
        with zipfile.ZipFile(path) as archive:
            return set(archive.namelist())

    def logs(self) -> list:
        directory = self.home / "logs"
        if not directory.exists():
            return []
        events = []
        for file in sorted(directory.glob("events-*.jsonl")):
            for line in file.read_text("utf-8").splitlines():
                events.append(json.loads(line))
        return events

    # -- value / structure ----------------------------------------------
    def test_set_cell_preserves_structure(self) -> None:
        before_parts = self.parts(self.rich)
        result = self.run_helper(
            self.rich, [{"op": "set_cell", "sheet": "数据", "cell": "B2", "value": "999", "expect": "200"}]
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("OK rich.xlsx", result.stdout)
        self.assertIn("校验：通过", result.stdout)

        workbook = load_workbook(self.rich)
        sheet = workbook["数据"]
        self.assertEqual(sheet["B2"].value, 999)
        self.assertEqual(sheet["B2"].number_format, "#,##0.00")
        self.assertEqual(sheet["C2"].value, "=B2/SUM($B$2:$B$5)")
        self.assertEqual(sheet["B6"].value, "=SUM(B2:B5)")
        self.assertEqual(sorted(sheet.merged_cells.ranges and [str(r) for r in sheet.merged_cells.ranges]),
                         ["A8:B8"])
        self.assertEqual(list(workbook.defined_names), ["基准"])
        self.assertTrue(sheet["C2"].comment is not None)
        self.assertEqual(len(sheet._charts), 1)
        self.assertEqual(len(sheet._images), 1)
        self.assertEqual(len(sheet.tables), 1)
        self.assertEqual(self.parts(self.rich) - before_parts, set())
        self.assertGreaterEqual(len(self.parts(self.rich)), len(before_parts) - 1)
        backups = list(self.rich.parent.glob("rich.bak-*.xlsx"))
        self.assertEqual(len(backups), 1)

    def test_expect_mismatch_keeps_file(self) -> None:
        digest = self.digest(self.rich)
        result = self.run_helper(
            self.rich, [{"op": "set_cell", "sheet": "数据", "cell": "B2", "value": "999", "expect": "123"}]
        )
        self.assertEqual(result.returncode, 2, result.stderr)
        self.assertIn("value_mismatch", result.stderr)
        self.assertEqual(self.digest(self.rich), digest)
        self.assertEqual(list(self.rich.parent.glob("rich.bak-*.xlsx")), [])

    def test_sheet_not_found(self) -> None:
        digest = self.digest(self.simple)
        result = self.run_helper(
            self.simple, [{"op": "set_cell", "sheet": "不存在", "cell": "A1", "value": "x"}]
        )
        self.assertEqual(result.returncode, 2, result.stderr)
        self.assertIn("not_found", result.stderr)
        self.assertIn("Sheet1", result.stderr)
        self.assertEqual(self.digest(self.simple), digest)

    def test_merged_non_top_left_refused(self) -> None:
        digest = self.digest(self.rich)
        result = self.run_helper(
            self.rich, [{"op": "set_cell", "sheet": "数据", "cell": "B8", "value": "x"}]
        )
        self.assertEqual(result.returncode, 3, result.stderr)
        self.assertIn("unsupported", result.stderr)
        self.assertIn("A8", result.stderr)
        self.assertEqual(self.digest(self.rich), digest)

    def test_formula_map_unchanged(self) -> None:
        result = self.run_helper(
            self.rich, [{"op": "set_cell", "sheet": "数据", "cell": "C4", "value": "0.5"}]
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        workbook = load_workbook(self.rich)
        sheet = workbook["数据"]
        self.assertEqual(sheet["C4"].value, 0.5)
        for row in range(2, 6):
            if row != 4:
                self.assertEqual(sheet[f"C{row}"].value, f"=B{row}/SUM($B$2:$B$5)")

    def test_value_coercion(self) -> None:
        result = self.run_helper(
            self.simple,
            [
                {"op": "set_cell", "sheet": "Sheet1", "cell": "A1", "value": "001"},
                {"op": "set_cell", "sheet": "Sheet1", "cell": "B2", "value": "12%"},
                {"op": "set_cell", "sheet": "Sheet1", "cell": "C4", "value": "=1+1"},
            ],
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        sheet = load_workbook(self.simple)["Sheet1"]
        self.assertEqual(sheet["A1"].value, "001")
        self.assertIsInstance(sheet["A1"].value, str)
        self.assertAlmostEqual(sheet["B2"].value, 0.12)
        self.assertEqual(sheet["C4"].value, "=1+1")

    # -- text replacement ------------------------------------------------
    def test_replace_text_all(self) -> None:
        result = self.run_helper(
            self.simple,
            [{"op": "replace_text", "sheet": "Sheet1", "find": "2025年度", "replace": "2026年度", "all": True}],
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        sheet = load_workbook(self.simple)["Sheet1"]
        self.assertEqual(sheet["B2"].value, "2026年度")
        self.assertEqual(sheet["B3"].value, "2026年度")

    def test_replace_text_ambiguous(self) -> None:
        digest = self.digest(self.simple)
        result = self.run_helper(
            self.simple,
            [{"op": "replace_text", "sheet": "Sheet1", "find": "2025年度", "replace": "2026年度"}],
        )
        self.assertEqual(result.returncode, 2, result.stderr)
        self.assertIn("ambiguous", result.stderr)
        self.assertEqual(self.digest(self.simple), digest)

    def test_replace_text_range(self) -> None:
        result = self.run_helper(
            self.simple,
            [{"op": "replace_text", "sheet": "Sheet1", "find": "2025年度", "replace": "2026年度",
              "range": "B2:B2"}],
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        sheet = load_workbook(self.simple)["Sheet1"]
        self.assertEqual(sheet["B2"].value, "2026年度")
        self.assertEqual(sheet["B3"].value, "2025年度")

    # -- formatting ------------------------------------------------------
    def test_set_format(self) -> None:
        result = self.run_helper(
            self.rich,
            [{"op": "set_format", "sheet": "数据", "cell": "B2", "font": "仿宋", "size": 12,
              "color": "FF0000", "bold": True, "number_format": "#,##0.00"}],
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        sheet = load_workbook(self.rich)["数据"]
        self.assertEqual(sheet["B2"].font.name, "仿宋")
        self.assertTrue(sheet["B2"].font.bold)
        self.assertEqual(str(sheet["B2"].font.color.rgb)[-6:], "FF0000")
        self.assertEqual(sheet["B2"].value, 200)

    # -- protection ------------------------------------------------------
    def test_protected_sheet_requires_confirmation(self) -> None:
        digest = self.digest(self.rich)
        result = self.run_helper(
            self.rich, [{"op": "set_cell", "sheet": "说明", "cell": "A2", "value": "x"}]
        )
        self.assertEqual(result.returncode, 3, result.stderr)
        self.assertIn("sheet_protected", result.stderr)
        self.assertEqual(self.digest(self.rich), digest)

        result = self.run_helper(
            self.rich, [{"op": "set_cell", "sheet": "说明", "cell": "A2", "value": "x"}], "--allow-protected"
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(load_workbook(self.rich)["说明"]["A2"].value, "x")

    # -- dry run / no-op -------------------------------------------------
    def test_dry_run(self) -> None:
        digest = self.digest(self.simple)
        result = self.run_helper(
            self.simple, [{"op": "set_cell", "sheet": "Sheet1", "cell": "C4", "value": "7"}], "--dry-run"
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("dry_run", result.stdout)
        self.assertEqual(self.digest(self.simple), digest)
        self.assertEqual(list(self.simple.parent.glob("simple.bak-*.xlsx")), [])

    def test_noop_write(self) -> None:
        digest = self.digest(self.simple)
        result = self.run_helper(
            self.simple, [{"op": "set_cell", "sheet": "Sheet1", "cell": "C4", "value": "100"}]
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("已是目标状态", result.stdout)
        self.assertEqual(self.digest(self.simple), digest)
        self.assertEqual(list(self.simple.parent.glob("simple.bak-*.xlsx")), [])

    # -- refusals --------------------------------------------------------
    def test_unknown_op(self) -> None:
        result = self.run_helper(self.simple, [{"op": "insert_row", "sheet": "Sheet1", "row": 3}])
        self.assertEqual(result.returncode, 3, result.stderr)
        self.assertIn("unsupported", result.stderr)

    def test_csv_refused(self) -> None:
        target = self.tmp / "table.csv"
        target.write_text("a,b\n1,2\n", "utf-8")
        result = self.run_helper(target, [{"op": "set_cell", "sheet": "Sheet1", "cell": "A1", "value": "x"}])
        self.assertEqual(result.returncode, 3, result.stderr)
        self.assertIn("unsupported", result.stderr)

    # -- warnings and logging -------------------------------------------
    def test_hidden_sheet_warning(self) -> None:
        result = self.run_helper(
            self.rich, [{"op": "set_cell", "sheet": "隐藏页", "cell": "A2", "value": "x"}]
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("隐藏", result.stdout)

    def test_logs_written(self) -> None:
        self.run_helper(self.simple, [{"op": "set_cell", "sheet": "Sheet1", "cell": "C4", "value": "7"}])
        events = self.logs()
        self.assertTrue(any(event.get("event") == "helper_ok" for event in events), events)
        ok = [event for event in events if event.get("event") == "helper_ok"][-1]
        self.assertEqual(ok.get("contract"), "v1.6")
        self.assertTrue(ok.get("ops"))
        self.assertEqual(ok["ops"][0].get("sheet"), "Sheet1")
        self.assertEqual(ok["ops"][0].get("cell"), "C4")

        self.run_helper(
            self.simple,
            [{"op": "replace_text", "sheet": "Sheet1", "find": "2025年度", "replace": "2026年度"}],
        )
        failures = [event for event in self.logs() if event.get("event") == "helper_fail"]
        self.assertTrue(failures)
        self.assertEqual(failures[-1].get("code"), "ambiguous")


if __name__ == "__main__":
    unittest.main(verbosity=2)
