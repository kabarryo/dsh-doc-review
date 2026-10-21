#!/usr/bin/env python3
"""Build fixture workbooks for the xlsx_edit.py test suite.

The rich fixture carries the structures a real ledger has: formulas, merged
cells, a table, data validation, conditional formatting, a chart, an image, a
comment, a hyperlink, a defined name, a protected sheet and a hidden sheet.
"""
from __future__ import annotations

import base64
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def build_rich(path: Path) -> Path:
    path = Path(path)
    png = path.with_name("dot.png")
    png.write_bytes(PNG)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet["A1"] = "项目"
    sheet["B1"] = "金额"
    sheet["C1"] = "占比"
    sheet["D1"] = "日期"
    for row in range(2, 6):
        sheet[f"A{row}"] = f"项目{row}"
        sheet[f"B{row}"] = row * 100
        sheet[f"C{row}"] = f"=B{row}/SUM($B$2:$B$5)"
        sheet[f"D{row}"] = datetime(2026, 10, row)
        sheet[f"D{row}"].number_format = "yyyy-mm-dd"
    sheet["B6"] = "=SUM(B2:B5)"
    sheet["B2"].number_format = "#,##0.00"
    sheet["A1"].font = Font(bold=True, name="黑体")
    sheet["A1"].fill = PatternFill("solid", fgColor="DDEBF7")
    sheet["E2"] = 0.125
    sheet["E2"].number_format = "0.00%"
    sheet.merge_cells("A8:B8")
    sheet["A8"] = "合并单元格"
    sheet.freeze_panes = "A2"
    sheet.print_area = "A1:D6"
    sheet.auto_filter.ref = "A1:D5"

    chart = BarChart()
    chart.title = "金额"
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=5))
    sheet.add_chart(chart, "F2")

    image = XLImage(str(png))
    sheet.add_image(image, "F20")

    sheet["C2"].comment = Comment("这是批注", "tester")
    sheet["A2"].hyperlink = "https://example.com"

    validation = DataValidation(type="whole", operator="between", formula1="0", formula2="1000")
    sheet.add_data_validation(validation)
    validation.add("B2:B5")

    sheet.conditional_formatting.add(
        "B2:B5",
        CellIsRule(operator="greaterThan", formula=["300"], fill=PatternFill("solid", fgColor="FFC7CE")),
    )

    table = Table(displayName="DataTbl", ref="A1:D5")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)

    note = workbook.create_sheet("说明")
    note["A1"] = "第二页"
    note.protection.sheet = True

    hidden = workbook.create_sheet("隐藏页")
    hidden["A1"] = "隐藏内容"
    hidden.sheet_state = "hidden"

    workbook.defined_names.add(DefinedName("基准", attr_text="数据!$B$2"))
    workbook.save(path)
    png.unlink(missing_ok=True)
    return path


def build_simple(path: Path) -> Path:
    path = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = "标题"
    sheet["B2"] = "2025年度"
    sheet["B3"] = "2025年度"
    sheet["C4"] = 100
    workbook.save(path)
    return path


if __name__ == "__main__":
    import sys

    target = Path(sys.argv[1] if len(sys.argv) > 1 else "fixtures")
    target.mkdir(parents=True, exist_ok=True)
    print(build_rich(target / "rich.xlsx"))
    print(build_simple(target / "simple.xlsx"))
