#!/usr/bin/env python3
"""Declarative DOCX edit helper for the dsh-doc-review plugin.

The plugin submits annotations; the agent turns them into a small ops list and
runs this helper once. The helper performs preflight, backup, matching, edit,
read-back verification and atomic replacement, and records failures in the
plugin log.

Usage:
  docx_edit.py --file /path/doc.docx --ops '[{"op":"replace","find":"a","replace":"b"}]'
  docx_edit.py --file /path/doc.docx --ops ops.json --dry-run
  docx_edit.py --file /path/doc.docx --ops ops.json --contract v1.4

Ops:
  {"op":"replace","find":"...","replace":"...", "all":false, "occurrence":1}
  {"op":"delete","find":"...", "all":false}
  {"op":"set_color","find":"...","color":"FF0000", "runs":"all"|"matched"}
  {"op":"set_font","find":"...","font":"楷体","size":12, "runs":"all"|"matched"}
  {"op":"set_paragraph_style","find":"...","style":"Heading 1"}

Scope: body paragraphs and top-level/nested tables. Headers, footers, text
boxes, fields, tracked changes and content controls are refused, not guessed.

Exit codes: 0 ok, 2 not-found/ambiguous, 3 unsupported, 4 verify failed,
5 file/exception.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import sys
import time
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
try:
    from doc_review_log import log_event as _log_event
except Exception:  # pragma: no cover - logger is optional at runtime
    def _log_event(**kwargs):  # type: ignore
        return None


def safe_log(**kwargs):
    """Logging must never fail or undo a successful edit."""
    try:
        return _log_event(**kwargs)
    except Exception:
        return None


class EditError(Exception):
    def __init__(self, code: str, message: str, stage: str = "edit", details: dict | None = None):
        super().__init__(message)
        self.code = code
        self.stage = stage
        self.details = details or {}


@dataclass
class Match:
    paragraph: object
    locator: str
    start: int
    end: int
    text: str


def sha256(path: Path) -> str:
    return "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()


def normalize(text: str) -> tuple[str, list[int]]:
    chars = []
    mapping = []
    for index, char in enumerate(text):
        if not char.isspace():
            chars.append(char)
            mapping.append(index)
    return "".join(chars), mapping


def iter_paragraphs(doc):
    for index, paragraph in enumerate(doc.paragraphs):
        yield paragraph, f"body[{index}]"
    for table_index, table in enumerate(doc.tables):
        yield from iter_table(table, f"table[{table_index}]")


def iter_table(table, prefix: str):
    for row_index, row in enumerate(table.rows):
        seen = set()
        for cell_index, cell in enumerate(row.cells):
            key = id(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            for paragraph_index, paragraph in enumerate(cell.paragraphs):
                yield paragraph, f"{prefix}.r{row_index}.c{cell_index}.p{paragraph_index}"
            for nested_index, nested in enumerate(cell.tables):
                yield from iter_table(nested, f"{prefix}.r{row_index}.c{cell_index}.t{nested_index}")


def paragraph_text(paragraph) -> str:
    return "".join(run.text for run in paragraph.runs)


def preflight(paragraph) -> str | None:
    xml = paragraph._p
    if xml.xpath(".//w:ins | .//w:del"):
        return "tracked_changes"
    if xml.xpath(".//w:fldSimple | .//w:instrText"):
        return "field"
    if xml.xpath(".//w:sdt"):
        return "content_control"
    if xml.xpath(".//w:txbxContent"):
        return "textbox"
    return None


def find_matches(doc, needle: str) -> list[Match]:
    target, _ = normalize(needle)
    if not target:
        raise EditError("bad_request", "find 不能为空", "preflight")
    matches: list[Match] = []
    for paragraph, locator in iter_paragraphs(doc):
        text = paragraph_text(paragraph)
        normalized, mapping = normalize(text)
        start = normalized.find(target)
        while start != -1:
            raw_start = mapping[start]
            raw_end = mapping[start + len(target) - 1] + 1
            matches.append(Match(paragraph, locator, raw_start, raw_end, text))
            start = normalized.find(target, start + len(target))
    return matches


def select_matches(matches: list[Match], op: dict) -> list[Match]:
    if not matches:
        raise EditError("not_found", f"未找到引文：{op.get('find')}", "match")
    if op.get("all"):
        return matches
    occurrence = op.get("occurrence")
    if occurrence:
        try:
            selected = matches[int(occurrence) - 1]
        except (IndexError, ValueError):
            raise EditError("not_found", f"occurrence={occurrence} 超出命中数 {len(matches)}", "match")
        return [selected]
    if len(matches) > 1:
        candidates = [{"locator": m.locator, "text": m.text[:100]} for m in matches[:8]]
        raise EditError(
            "ambiguous",
            f"引文命中 {len(matches)} 处，未指定 all/occurrence",
            "match",
            {"candidates": candidates},
        )
    return matches


def run_span(paragraph, start: int, end: int):
    runs = paragraph.runs
    position = 0
    first = None
    last = None
    for index, run in enumerate(runs):
        length = len(run.text)
        if first is None and start < position + length:
            first = (index, start - position)
        if first is not None and end <= position + length:
            last = (index, end - position)
            break
        position += length
    return first, last


def apply_replace(match: Match, replacement: str) -> str:
    paragraph = match.paragraph
    first, last = run_span(paragraph, match.start, match.end)
    if first is None or last is None:
        raise EditError("unsupported", "引文跨越了不可编辑的 run 结构", "edit")
    first_index, first_offset = first
    last_index, last_offset = last
    runs = paragraph.runs
    if first_index == last_index:
        run = runs[first_index]
        run.text = run.text[:first_offset] + replacement + run.text[last_offset:]
    else:
        head = runs[first_index]
        tail = runs[last_index]
        head.text = head.text[:first_offset] + replacement
        for index in range(first_index + 1, last_index):
            runs[index].text = ""
        tail.text = tail.text[last_offset:]
    return paragraph_text(paragraph)


def style_runs(match: Match, op: dict):
    runs = match.paragraph.runs
    if op.get("runs", "all") == "matched":
        first, last = run_span(match.paragraph, match.start, match.end)
        if first is None or last is None:
            return []
        return runs[first[0] : last[0] + 1]
    return list(runs)


def set_color(run, color: str) -> None:
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    node = rpr.find(qn("w:color"))
    if node is not None:
        for attr in ("w:themeColor", "w:themeTint", "w:themeShade"):
            node.attrib.pop(qn(attr), None)


def set_font(run, font: str, size=None) -> None:
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(float(size))


def apply_op(doc, op: dict) -> dict:
    kind = op.get("op") or op.get("type")
    if kind not in {"replace", "delete", "set_color", "set_font", "set_paragraph_style"}:
        raise EditError("unsupported", f"不支持的 op：{kind}", "preflight")
    find = op.get("find") or ""
    matches = find_matches(doc, find)
    selected = select_matches(matches, op)
    for match in selected:
        reason = preflight(match.paragraph)
        if reason:
            raise EditError("unsupported", f"目标段落含不支持的结构：{reason}", "preflight", {"locator": match.locator})

    if kind in {"replace", "delete"}:
        replacement = "" if kind == "delete" else str(op.get("replace", ""))
        # Apply right-to-left per paragraph so offsets stay valid with all=true.
        ordered = sorted(selected, key=lambda m: (m.locator, -m.start))
        changed = []
        for match in ordered:
            changed.append({"locator": match.locator, "before": match.text[match.start : match.end], "after": replacement})
            apply_replace(match, replacement)
        return {"op": kind, "status": "ok", "changed": changed}

    if kind == "set_paragraph_style":
        style = op.get("style")
        for match in selected:
            try:
                match.paragraph.style = style
            except Exception as exc:
                raise EditError("unsupported", f"样式不存在：{style}", "edit", {"error": str(exc)})
        return {"op": kind, "status": "ok", "changed": [{"locator": m.locator, "style": style} for m in selected]}

    changed = []
    for match in selected:
        runs = style_runs(match, op)
        if not runs:
            continue
        for run in runs:
            if kind == "set_color":
                set_color(run, str(op.get("color", "FF0000")).lstrip("#").upper())
            else:
                set_font(run, str(op.get("font", "仿宋_GB2312")), op.get("size"))
        changed.append({"locator": match.locator, "runs": len(runs)})
    return {"op": kind, "status": "ok", "changed": changed}


def verify(doc_path: Path, ops: list[dict], report: list[dict]) -> None:
    doc = Document(str(doc_path))
    for op, result in zip(ops, report):
        kind = op.get("op") or op.get("type")
        find = op.get("find") or ""
        matches = find_matches(doc, find)
        if kind in {"replace", "delete"}:
            replacement = "" if kind == "delete" else str(op.get("replace", ""))
            if replacement:
                if not find_matches(doc, replacement):
                    raise EditError("verify_failed", f"回读未找到新文案：{replacement}", "verify")
            elif find_matches(doc, find):
                raise EditError("verify_failed", f"删除后仍能匹配旧文案：{find}", "verify")
        elif kind == "set_color":
            color = str(op.get("color", "FF0000")).lstrip("#").upper()
            for match in matches:
                for run in style_runs(match, op):
                    current = run.font.color.rgb
                    if current is None or str(current).upper() != color:
                        raise EditError("verify_failed", f"颜色未生效：{match.locator}", "verify")
        elif kind == "set_font":
            font = str(op.get("font", ""))
            for match in matches:
                for run in style_runs(match, op):
                    if run.font.name != font:
                        raise EditError("verify_failed", f"字体未生效：{match.locator}", "verify")
        elif kind == "set_paragraph_style":
            for match in matches:
                if match.paragraph.style.name != op.get("style"):
                    raise EditError("verify_failed", f"段落样式未生效：{match.locator}", "verify")


def zip_parts(path: Path) -> dict:
    with zipfile.ZipFile(path) as archive:
        return {name: hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist()}


def backup_file(path: Path, digest: str) -> Path:
    for candidate in sorted(path.parent.glob(f"{path.stem}.bak-*.docx")):
        try:
            if sha256(candidate) == digest:
                return candidate
        except OSError:
            continue
    backup = path.with_name(f"{path.stem}.bak-{datetime.now().strftime('%Y%m%d-%H%M%S')}.docx")
    shutil.copy2(path, backup)
    return backup


def load_ops(raw: str) -> list[dict]:
    path = Path(raw).expanduser()
    if path.is_file():
        raw = path.read_text("utf-8")
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise EditError("bad_request", f"ops JSON 解析失败：{exc}", "preflight")
    if isinstance(data, dict):
        data = data.get("operations") or data.get("ops") or [data]
    if not isinstance(data, list) or not data:
        raise EditError("bad_request", "ops 必须是非空数组", "preflight")
    return data


def main() -> None:
    parser = argparse.ArgumentParser(description="docx edit helper")
    parser.add_argument("--file", required=True)
    parser.add_argument("--ops", required=True, help="ops JSON file or inline JSON")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--contract")
    parser.add_argument("--json", action="store_true", dest="as_json")
    args = parser.parse_args()

    started = time.monotonic()
    path = Path(args.file).expanduser().resolve()
    contract = args.contract or "v1.5"
    try:
        if not path.is_file():
            raise EditError("file_error", f"文件不存在：{path}", "preflight")
        ops = load_ops(args.ops)
        doc = Document(str(path))
        before_digest = sha256(path)
        before_parts = zip_parts(path)
        report = [apply_op(doc, op) for op in ops]
        after_texts = {}
        if args.dry_run:
            print(json.dumps({"ok": True, "dry_run": True, "file": str(path), "ops": report}, ensure_ascii=False, indent=2))
            return
        backup = backup_file(path, before_digest)
        temp = path.with_name(path.stem + ".dsh-tmp" + path.suffix)
        doc.save(str(temp))
        verify(temp, ops, report)
        after_parts = zip_parts(temp)
        changed_parts = [name for name, digest in after_parts.items() if before_parts.get(name) != digest]
        unexpected = [name for name in changed_parts if name != "word/document.xml"]
        os.chmod(temp, path.stat().st_mode)
        os.replace(temp, path)
        after_digest = sha256(path)
        result = {
            "ok": True,
            "file": str(path),
            "backup": str(backup),
            "hash_before": before_digest,
            "hash_after": after_digest,
            "ops": report,
            "verification": {"ok": True, "changed_parts": changed_parts, "unexpected_parts": unexpected},
        }
        safe_log(
            event="helper_ok",
            contract=contract,
            file=str(path),
            file_hash=before_digest,
            duration_ms=int((time.monotonic() - started) * 1000),
            ops=ops,
        )
        if args.as_json:
            print(json.dumps(result, ensure_ascii=False, indent=2))
        else:
            print(f"OK {path.name}")
            print(f"备份：{backup.name}")
            print(f"hash：{before_digest[7:19]}… -> {after_digest[7:19]}…")
            for index, item in enumerate(report, 1):
                kind = item["op"]
                changes = item.get("changed") or []
                if kind in {"replace", "delete"}:
                    for change in changes:
                        before = (change.get("before") or "")[:30]
                        after = (change.get("after") or "")[:30]
                        print(f"{index}. {kind} {change.get('locator')}: “{before}” -> “{after}”")
                elif kind == "set_paragraph_style":
                    for change in changes:
                        print(f"{index}. {kind} {change.get('locator')}: {change.get('style')}")
                else:
                    detail = ops[index - 1].get("font") or ops[index - 1].get("color") or ""
                    for change in changes:
                        print(f"{index}. {kind} {change.get('locator')}: {detail} ({change.get('runs')} runs)")
            print(f"校验：通过；变更部件：{', '.join(changed_parts) or '无'}")
            if unexpected:
                print(f"警告：额外变更部件：{', '.join(unexpected)}")
    except EditError as exc:
        safe_log(
            event="verify_fail" if exc.code == "verify_failed" else "helper_fail",
            contract=contract,
            file=str(path),
            stage=exc.stage,
            code=exc.code,
            reason=str(exc),
            duration_ms=int((time.monotonic() - started) * 1000),
            ops=ops if "ops" in locals() else None,
            note=json.dumps(exc.details, ensure_ascii=False) if exc.details else None,
        )
        print(f"ERROR [{exc.code}] {exc}", file=sys.stderr)
        if exc.details:
            print(json.dumps(exc.details, ensure_ascii=False), file=sys.stderr)
        raise SystemExit(4 if exc.code == "verify_failed" else 3 if exc.code == "unsupported" else 2 if exc.code in {"not_found", "ambiguous", "bad_request"} else 5)
    except Exception as exc:  # noqa: BLE001
        safe_log(
            event="helper_fail",
            contract=contract,
            file=str(path),
            stage="exception",
            code="exception",
            reason=f"{type(exc).__name__}: {exc}",
            duration_ms=int((time.monotonic() - started) * 1000),
            ops=ops if "ops" in locals() else None,
        )
        print(f"ERROR [exception] {type(exc).__name__}: {exc}", file=sys.stderr)
        raise SystemExit(5)


if __name__ == "__main__":
    main()
